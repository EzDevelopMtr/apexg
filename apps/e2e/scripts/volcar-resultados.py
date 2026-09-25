"""Vuelca el resultado de Playwright en el plan de pruebas.

Lee reports/results.json y llena las columnas amarillas de
docs/pruebas/plan-pruebas-e2e.xlsx. La fila se ubica por el id del caso
(CP-xx), que Playwright arrastra en el titulo de cada prueba.
"""

import json
import os
import re
import sys
from datetime import date

from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
RESULTADOS = os.path.join(RAIZ, "apps", "e2e", "reports", "results.json")
PLAN = os.path.join(RAIZ, "docs", "pruebas", "plan-pruebas-e2e.xlsx")

ESTADO = {
    "passed": "OK",
    "failed": "Fallo",
    "timedOut": "Fallo",
    "interrupted": "Bloqueado",
    "skipped": "No aplica",
}


def recorrer(nodo):
    for suite in nodo.get("suites", []):
        yield from recorrer(suite)
    for spec in nodo.get("specs", []):
        yield spec


def limpiar(texto):
    return re.sub(r"\x1b\[[0-9;]*m", "", texto or "").strip()


def main():
    with open(RESULTADOS, encoding="utf-8") as archivo:
        reporte = json.load(archivo)

    por_caso = {}
    for suite in reporte.get("suites", []):
        for spec in recorrer(suite):
            caso = re.match(r"(CP-\d+)", spec.get("title", ""))
            if not caso:
                continue
            prueba = (spec.get("tests") or [{}])[0]
            resultado = (prueba.get("results") or [{}])[0]
            estado = ESTADO.get(resultado.get("status"), "Pendiente")
            mensaje = limpiar((resultado.get("error") or {}).get("message", ""))
            evidencias = [
                os.path.basename(adjunto.get("path") or adjunto.get("name", ""))
                for adjunto in resultado.get("attachments", [])
            ]
            por_caso[caso.group(1)] = {
                "estado": estado,
                "obtenido": mensaje[:500] if mensaje else "Comportamiento conforme a lo esperado.",
                "evidencia": ", ".join(e for e in evidencias if e)[:250]
                or "reports/html (traza y video)",
                "duracion": resultado.get("duration", 0),
            }

    if not por_caso:
        print("No se encontro ningun caso CP-xx en el reporte.", file=sys.stderr)
        return 1

    libro = load_workbook(PLAN)
    hoja = libro["Casos de prueba"]
    hoy = date.today().isoformat()
    escritos = 0

    for fila in range(2, hoja.max_row + 1):
        identificador = hoja.cell(row=fila, column=1).value
        dato = por_caso.get(identificador)
        if not dato:
            continue
        hoja.cell(row=fila, column=11).value = dato["estado"]
        hoja.cell(row=fila, column=12).value = dato["obtenido"]
        hoja.cell(row=fila, column=13).value = dato["evidencia"]
        hoja.cell(row=fila, column=14).value = (
            f"Automatizado con Playwright ({dato['duracion'] / 1000:.1f} s)."
        )
        hoja.cell(row=fila, column=15).value = hoy
        escritos += 1

    libro.calculation.fullCalcOnLoad = True
    libro.save(PLAN)
    print(f"{escritos} casos volcados de {len(por_caso)} ejecutados.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
